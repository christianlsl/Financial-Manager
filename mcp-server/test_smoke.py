"""Smoke test for the MCP server tools against a live backend.

Requires:
- backend running on 127.0.0.1:9910
- env vars FM_TEST_EMAIL / FM_TEST_PASSWORD set to a real account
  (a second account FM_TEST_EMAIL2 / FM_TEST_PASSWORD2 is optional,
  used for multi-user isolation checks)
- FM_API_BASE_URL (optional, defaults to http://127.0.0.1:9910)

Run: uv run python test_smoke.py
"""

import json
import os
import sys
import base64
from datetime import date, timedelta
from io import BytesIO

os.environ.setdefault("FM_API_BASE_URL", "http://127.0.0.1:9910")
EMAIL = os.environ.get("FM_TEST_EMAIL", "")
PASSWORD = os.environ.get("FM_TEST_PASSWORD", "")
EMAIL2 = os.environ.get("FM_TEST_EMAIL2", "")
PASSWORD2 = os.environ.get("FM_TEST_PASSWORD2", "")
if not EMAIL or not PASSWORD:
    print("请先设置 FM_TEST_EMAIL / FM_TEST_PASSWORD 环境变量")
    sys.exit(1)
os.environ.setdefault("FM_INVOICE_DIR", "/tmp/fm-mcp-invoices")

from mcp_server import api_client  # noqa: E402
from mcp_server.api_client import NotLoggedInError  # noqa: E402
from mcp_server.tools import auth, invoices, mutations, queries  # noqa: E402


class FakeCtx:
    """Minimal stand-in for FastMCP Context: only .session is needed."""

    def __init__(self) -> None:
        self.session = object()


ok = 0


def check(name, cond, extra=""):
    global ok
    if cond:
        ok += 1
        print(f"[PASS] {name} {extra}")
    else:
        print(f"[FAIL] {name} {extra}")
        sys.exit(1)


def main():
    ctx = FakeCtx()
    other = FakeCtx()  # a second, not-logged-in session

    # --- not logged in yet ---
    try:
        queries.get_summary(ctx)
        check("blocked before login", False)
    except NotLoggedInError as e:
        check("blocked before login", "登录" in str(e))

    # --- wrong password ---
    try:
        auth.login(EMAIL, "definitely-wrong-password", ctx)
        check("wrong password rejected", False)
    except api_client.BackendError:
        check("wrong password rejected", True)
    except NotLoggedInError:
        check("wrong password rejected", False, "got NotLoggedInError instead")

    # --- login via tool ---
    welcome = auth.login(EMAIL, PASSWORD, ctx)
    check("login", "登录成功" in welcome, welcome[:60])

    # --- whoami / logout / re-login ---
    me = json.loads(auth.whoami(ctx))
    check("whoami", me.get("email") == EMAIL, f"email={me.get('email')}")
    auth.login(EMAIL, PASSWORD, ctx)  # re-login overwrites cleanly

    # --- queries ---
    s = json.loads(queries.get_summary(ctx))
    check("get_summary", "monthly" in s)

    sales = json.loads(queries.list_sales(ctx=ctx, limit=5))
    check("list_sales", "items" in sales and "total" in sales)
    print("   sales total:", sales["total"])

    customers = json.loads(queries.list_customers(ctx=ctx))
    check("list_customers", isinstance(customers, list))

    suppliers = json.loads(queries.list_suppliers(ctx=ctx))
    check("list_suppliers", isinstance(suppliers, list))

    types = json.loads(queries.list_types(ctx))
    check("list_types", isinstance(types, list))

    stats = json.loads(queries.get_statistics(ctx=ctx, analysis_type="monthly"))
    check("get_statistics", "overview" in stats or isinstance(stats, dict))

    # --- session isolation: another session is still blocked ---
    try:
        queries.list_sales(ctx=other)
        check("other session isolated", False)
    except NotLoggedInError:
        check("other session isolated", True)

    # --- optional: two users see different data ---
    if EMAIL2 and PASSWORD2:
        auth.login(EMAIL2, PASSWORD2, other)
        sales2 = json.loads(queries.list_sales(ctx=other, limit=5))
        check(
            "two users both work",
            "items" in sales2,
            f"user2 email={EMAIL2} total={sales2.get('total')}",
        )
        auth.logout(other)
        try:
            queries.list_sales(ctx=other)
            check("logout clears session", False)
        except NotLoggedInError:
            check("logout clears session", True)

    # --- create/update/delete sale ---
    created = json.loads(mutations.create_sale(
        date=date.today(),
        item_name="冒烟测试商品",
        items_count=3,
        unit_price=25.5,
        notes="MCP smoke test",
        ctx=ctx,
    ))
    sale_id = created["id"]
    check("create_sale", sale_id > 0, f"id={sale_id} total={created['total_price']}")

    updated = json.loads(mutations.update_sale(
        sale_id=sale_id,
        status="sent",
        notes="updated by smoke test",
        ctx=ctx,
    ))
    check("update_sale", updated["status"] == "sent", f"status={updated['status']}")

    fetched = json.loads(queries.get_sale(sale_id, ctx=ctx))
    check("get_sale", fetched["id"] == sale_id and fetched["total_price"] == "76.50")

    # --- batch invoices ---
    from mcp.types import CallToolResult, EmbeddedResource, TextContent
    batch = invoices.generate_invoices_batch(
        date_from=(date.today() - timedelta(days=1)).isoformat(),
        date_to=date.today().isoformat(),
        ctx=ctx,
    )
    check("batch invoices", isinstance(batch, CallToolResult))
    embedded = [r for r in batch.content if isinstance(r, EmbeddedResource)]
    texts = [r for r in batch.content if isinstance(r, TextContent)]
    check("batch has xlsx attachment", len(embedded) == 1, f"blob len={len(embedded[0].resource.blob)}")
    check("batch has summary", len(texts) == 1)

    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(base64.b64decode(embedded[0].resource.blob)))
    ws = wb.active
    headers = [ws.cell(row=1, column=i).value for i in range(1, 13)]
    expected_headers = ["日期", "项目", "公司", "部门", "客户", "类型", "数量", "单价", "金额", "图片", "状态", "备注"]
    check("batch headers", headers == expected_headers, f"headers={headers}")
    check("batch has data rows", ws.max_row >= 2, f"rows={ws.max_row}")
    wb.close()

    companies = json.loads(queries.list_companies(ctx=ctx))
    if companies:
        company_batch = invoices.generate_invoices_batch(company_id=companies[0]["id"], limit=5, ctx=ctx)
        check("batch company_id call", company_batch is not None)

    # --- cleanup ---
    deleted = json.loads(mutations.delete_sale(sale_id, ctx=ctx))
    check("delete_sale", "ok" in str(deleted).lower() or deleted is not None)

    auth.logout(ctx)
    try:
        queries.get_summary(ctx)
        check("logout clears own session", False)
    except NotLoggedInError:
        check("logout clears own session", True)

    print(f"\nAll {ok} checks passed.")


if __name__ == "__main__":
    main()
